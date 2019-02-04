import os
import base64
import tempfile
import uuid

from fpdf import FPDF
from docx import Document
from odoo import models, fields, api, http
from odoo.exceptions import UserError
from PyPDF2 import PdfFileReader, PdfFileWriter
from odoo.addons.dn_base.statics import raise_dn_model_error

from odoo.tools import pycompat, datetime, DEFAULT_SERVER_DATETIME_FORMAT


class MeetingDoc(models.Model):
    _name = 'meeting_point.doc'
    _inherit = 'dn_documents.allfiles'
    meeting_id = fields.Many2one('calendar.event', string="Meeting", ondelete='cascade')

    @api.model
    def search(self, args, offset=0, limit=0, order=None, count=False):
        if not self.env.user.has_group('dn_base.group_dn_app_manager') and not self.env.user.has_group('base.group_system'):
            # list = self.env['calendar.event'].search([('partner_ids', 'in', [self.env.user.partner_id.id])]).ids
            myargs = [('meeting_id.partner_ids','in',[self.env.user.partner_id.id])]
            args.extend(myargs)
        docs = super(MeetingDoc, self).search(args)
        return docs

    def get_audience(self):
        ids = []
        for partner in self.meeting_id.partner_ids:
            ids.append(partner.user_id.id)
        return ids

class Document(models.Model):
    _name = 'meeting_point.document'
    _inherit = ['e_sign.document']#,'dn.seen'

    meeting_id = fields.Many2one('calendar.event', string="Meeting", ondelete='cascade')
    user_ids = fields.Many2many('res.users', string="Signature(s)")
    # seen_by_me = fields.Integer(compute='_compute_seen_by_me', default=0)
    mp_signature_status = fields.Char(string='My Signature', compute="_compute_signature_status")

    @api.multi
    def _compute_signature_status(self):
        for doc in self:
            pending = False
            found = False
            if doc.meeting_id and not doc.meeting_id.publish:
                doc.mp_signature_status = "Unpublished"
                continue
            doc.mp_signature_status = "Not required"
            for signature in doc.signature_ids:
                if signature.user_id == self.env.user:
                    found = True
                    if not signature.draw_signature:
                        pending = True
            if pending:
                doc.mp_signature_status ="Pending"
            else:
                if found:
                    doc.mp_signature_status = "Completed"


    @api.onchange('name')
    def filter_users(self):
        users=[]
        for p in self.meeting_id.partner_ids:
            if p.user_id:
                users.append(p.user_id.id)
        domain = {'domain': {'user_ids': [('id', 'in', users)]}}
        return domain

    @api.model
    def create(self, values):
        try:
            #user_ids = values['user_ids'][0][2]
            doc = super(Document, self).create(values)
            user_ids = doc.user_ids
            if user_ids and len(user_ids) > 0:
                for user in user_ids:
                    self.create_signature(user.id, doc)
                self.sudo().embed_signature(doc)
            return doc
        except:
            raise_dn_model_error("Error while creating document \n")

    def write(self, values):
        try:
            before_ids = self.user_ids
            super(Document, self).write(values)
            if before_ids != self.user_ids:
                self.update_signatures()
            return True
        except:
            raise_dn_model_error("Error while writing document \n")

    def create_signature(self, uid, doc):
        token = pycompat.text_type(uuid.uuid4())
        obj = { 'user_id': uid,'email': '','type': 'sign','token': token,'document_id': doc.id}
        doc.signature_ids = [(0, 0, obj)]
        sign=doc.signature_ids.filtered(lambda r: r.token==token)
        doc.send_mail(uid,doc,sign.id)


    def send_mail(self,uid,doc,sign_id):

        self = self.with_context(doc_name=self.name)
        invitation_template = self.env.ref('meeting_point.email_template_signature_mp')
        user=self.env['res.users'].search([('id','=',uid)])
        email_values={"recipient_ids":[(4, user.partner_id.id)]}
        invitation_template.send_mail(sign_id,force_send=True, raise_exception=False, email_values=email_values)

    def notify_signed_mail(self,sign_id):
        self = self.with_context(doc_name=self.name)
        invitation_template = self.env.ref('meeting_point.email_template_signed_notify_mp')
        invitation_template.send_mail(sign_id,force_send=True, raise_exception=False)

    def update_signatures(self):
        for doc in self:
            #Delete sign of removed authorities
            user_ids = self.user_ids
            for sign in doc.signature_ids:
                found = False
                for user in user_ids:
                    if sign.user_id.id == user.id:
                        found = True
                if not found:
                    sign.unlink()

            # Insert sign for new authorities
            for user in user_ids:
                found = False
                for sign in doc.signature_ids:
                    if sign.user_id.id == user.id:
                        found = True
                if not found:
                    self.create_signature(user.id, doc)
                self.sudo().embed_signature(doc)

    def embed_signature(self, doc, ip=None):
        if not doc.original_pdf or len(doc.signature_ids) == 0:
            return
        #pth = tempfile.gettempdir()
        if not ip:
            ip = http.request.httprequest.remote_addr
        curr_dir = os.path.dirname(__file__)
        pth = curr_dir.replace('models', 'doc_signs')

        file = doc.original_pdf
        f = base64.b64decode(file)
        # f = base64.decodestring(file)

        fobj = tempfile.NamedTemporaryFile(delete=False)
        fname1 = fobj.name
        fobj.write(f)
        fobj.close()

        input = PdfFileReader(open(fname1, "rb"))
        # Addition of code for orientation correction Asfand
        pageValue = input.getPage(0)
        pageOrientation = pageValue.get('/Rotate')
        page = input.getPage(0).mediaBox
        zAxis = page.getUpperRight_x()
        yAxis = page.getUpperLeft_x()
        width = int(zAxis - yAxis)

        height = int(page.getUpperRight_y() - page.getLowerRight_y())
        solution = [width, height]
        if width > height or pageOrientation == 90:
            if height > width:
                orientation = 'L'
            else:
                orientation = 'P'
        elif pageOrientation == 0 or pageOrientation == 180 or pageOrientation == None:
            if width > height:
               orientation = 'L'
            else:
                orientation = 'P'

        output = PdfFileWriter()
        count = 0
        pdf = FPDF(orientation, 'pt', solution)
        pdf.add_page(orientation=orientation)
        # Addition ended

        for page_number in range(input.getNumPages()):
            output.addPage(input.getPage(page_number))

        for sign in doc.signature_ids:
            pdf.ln(60)
            if sign.draw_signature:
                file_s = sign.draw_signature
                # f = base64.decodestring(file_s)
                f = base64.b64decode(file_s)
                # create a writable image and write the decoding result
                exten = 'png'
                if sign.filename:
                    exten = self.get_extension(sign.filename)

                signarure_image_path = pth + "/signature-" + str(sign.id) + "." + exten
                #signarure_image_path = "r"+"'"+ signarure_image_path+"'" =>gives error
                image_result = open(signarure_image_path,'wb')
                image_result.write(f)
                image_result.close()
                pdf.image(signarure_image_path, x=-2, y=None, w=345, h=150)
            pdf.ln(5)
            pdf.set_font('Arial', 'U', 15)
            signature_authority = sign.user_id
            # Special case, we have restricted our users to read super admin
            try:
                sign_name = signature_authority.name
            except:
                sign_name = 'Root'
            pdf.cell(5, 5, sign_name)
            if sign.draw_signature:
                date = datetime.strptime(sign.write_date,
                                  DEFAULT_SERVER_DATETIME_FORMAT)
                date = date.strftime('%b %d %Y')
                pdf.ln(20)
                pdf.cell(5, 5, date)
                if ip:
                    pdf.ln(25)
                    pdf.cell(5, 5, ip)
            count = count + 1


            if (count == 9):
                pdf.add_page(orientation=orientation)
                count = 0


        signature_only_pdf_path = pth + "/signature-pdf-"+str(doc.id)+".pdf"
        try:
            pdf.output(signature_only_pdf_path, "F")
        except:
            if not os.path.exists(pth):
                os.makedirs(pth)
                pdf.output(signature_only_pdf_path, "F")
            else:
                raise
        pdf.close()

        signaturepdf = PdfFileReader(open(signature_only_pdf_path, "rb"))
        for page_number in range(signaturepdf.getNumPages()):
            output.addPage(signaturepdf.getPage(page_number))

        output_pdf_path = pth + "/signed-doc-output-"+str(doc.id)+".pdf"
        with open(output_pdf_path, "wb") as outputStream:
            output.write(outputStream)
        res = open(output_pdf_path, 'rb')
        read = res.read()
        # res_encode = base64.encodestring(read)
        res_encode = base64.b64encode(read)
        res.close()

        self.sudo().write({'pdf_doc': res_encode, 'from_code': True})

    def get_extension(self, filename):
        ar = filename.split('.')
        return ar[len(ar)-1]

    def open_signature_form(self):

        view_id = self.env.ref('meeting_point.view_meeting_doc_view_sign_form').id
        if self:
            return {
                'type': 'ir.actions.act_window',
                'name': self.name,
                'view_id': view_id,
                'view_mode': 'form',
                'res_model': self._name,
                'res_id': self.id,
                'target': 'new',
            }

    def assign_signatures(self):
        view_id = self.env.ref('meeting_point.view_meeting_doc_edit_form').id
        return {
            'type': 'ir.actions.act_window',
            'name': 'Assign Signatures',
            'view_id': view_id,
            'view_mode': 'form',
            'res_model': self._name,
            'res_id': self.id,
            'target': 'current',
        }

    @api.model
    def search(self, args, offset=0, limit=0, order=None, count=False):
        if not self.env.user.has_group('dn_base.group_dn_app_manager') and not self.env.user.has_group('base.group_system'):
            list = self.env['e_sign.signature'].search([('user_id', '=', self.env.user.id)]).ids
            myargs = [('signature_ids','in',list)]
            args.extend(myargs)
        docs = super(Document, self).search(args)
        return docs
